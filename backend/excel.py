"""Regenerate a separate workbook from the committed ledger, never the source."""
import os
import tempfile
import threading
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.properties import CalcProperties

from .domain import TICKET_CAPACITY, now_iso
from .reports import audit_reports


def safe_text(value):
    """Do not turn guest names, notes or email data into spreadsheet formulas."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


class ExcelExporter:
    def __init__(self, store, output):
        self.store = store
        self.output = Path(output)
        self.output.parent.mkdir(parents=True, exist_ok=True)
        self.lock = threading.Lock()

    def sync(self):
        if not self.lock.acquire(blocking=False):
            return False
        temporary = None
        book = None
        try:
            # Hold the store lock through replacement so an old snapshot cannot mark a newer revision exported.
            with self.store.snapshot():
                locks = [self.output.with_name("~$" + self.output.name), self.output.with_name("~$" + self.output.name[2:])]
                if any(path.exists() for path in locks):
                    raise PermissionError("Close the live workbook in Excel to allow updates.")
                book = load_workbook(self.store.template)
                sheet = book["Raffle List"]
                guests = self.store.guests()
                payments = self.store.payments()
                references = {}
                for payment in payments:
                    if payment["status"] == "confirmed":
                        for allocation in payment["allocations"]:
                            references.setdefault(allocation["guest_id"], set()).add(payment["id"])
                for guest in guests:
                    row = guest["source_row"]
                    sheet.cell(row, 6, guest["ticket_count"])
                    sheet.cell(row, 7, guest["paid_cents"] / 100).number_format = '"$"#,##0.00'
                    sheet.cell(row, 8, ", ".join(sorted(references.get(guest["id"], []))))
                summary = self.store.summary()
                sheet["K5"] = TICKET_CAPACITY
                last_row = max(g["source_row"] for g in guests)
                sheet["K6"] = f"=SUM(F5:F{last_row})"
                sheet["K7"] = f'=K5-K6-COUNTIF(\'Ticket Register\'!B2:B1501,"void")'
                sheet["K8"] = f"=SUM(G5:G{last_row})"
                sheet["J7"] = "Unissued tickets (excludes voids)"
                sheet.auto_filter.ref = f"B4:H{last_row}"
                for report in audit_reports(self.store)["sheets"]:
                    tab = book.create_sheet(report["name"])
                    tab.append([column["label"] for column in report["columns"]])
                    for row in report["rows"]:
                        tab.append([value / 100 if column["type"] == "money" and value is not None else safe_text(value)
                                    for value, column in zip(row, report["columns"])])
                    tab.freeze_panes = "A2"
                    tab.auto_filter.ref = tab.dimensions
                    for cell in tab[1]:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = PatternFill("solid", fgColor="65715E")
                    for column in tab.columns:
                        letter = column[0].column_letter
                        width = min(44, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
                        tab.column_dimensions[letter].width = width
                    for index, column in enumerate(report["columns"], start=1):
                        if column["type"] == "money":
                            for row in tab.iter_rows(min_row=2, min_col=index, max_col=index):
                                row[0].number_format = '\"$\"#,##0.00'
                book.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
                descriptor, temporary = tempfile.mkstemp(suffix=".xlsx", dir=self.output.parent)
                os.close(descriptor)
                book.save(temporary)
                book.close()
                book = None
                os.chmod(temporary, 0o600)
                os.replace(temporary, self.output)
                self.store.set_values(excel_pending=False, excel_error="", last_export=now_iso(),
                                      exported_revision=summary["revision"])
                return True
        except Exception as exc:
            self.store.set_values(excel_pending=True, excel_error=f"Excel update pending: {type(exc).__name__}. Close the live workbook and retry.")
            return False
        finally:
            if book is not None:
                book.close()
            if temporary and os.path.exists(temporary):
                os.unlink(temporary)
            self.lock.release()
