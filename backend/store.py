import hashlib
import json
import os
import shutil
import sqlite3
import threading
import uuid
from contextlib import contextmanager
from pathlib import Path

from openpyxl import load_workbook

from .domain import FIRST_TICKET, LAST_TICKET, TICKET_CAPACITY, RuleError, now_iso, prize_cents, ticket_count


class Store:
    def __init__(self, data_dir: Path, source: Path):
        self.root = Path(data_dir)
        self.root.mkdir(parents=True, exist_ok=True, mode=0o700)
        os.chmod(self.root, 0o700)
        self.path = self.root / "raffle.sqlite3"
        self.template = self.root / "registration-template.xlsx"
        self.lock = threading.RLock()
        with self.connection() as db:
            db.executescript("""
                PRAGMA journal_mode=WAL;
                CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS guests(
                    id TEXT PRIMARY KEY, name TEXT NOT NULL, remarks TEXT, guest_of TEXT,
                    table_no TEXT, source_row INTEGER NOT NULL UNIQUE);
                CREATE TABLE IF NOT EXISTS payments(
                    id TEXT PRIMARY KEY, source_key TEXT UNIQUE, source TEXT NOT NULL,
                    amount_cents INTEGER NOT NULL CHECK(amount_cents>0), bank_time TEXT,
                    account_suffix TEXT, created_at TEXT NOT NULL, confirmed_at TEXT,
                    status TEXT NOT NULL DEFAULT 'pending', note TEXT NOT NULL DEFAULT '',
                    source_link TEXT NOT NULL DEFAULT '');
                CREATE TABLE IF NOT EXISTS allocations(
                    id TEXT PRIMARY KEY, payment_id TEXT NOT NULL REFERENCES payments(id),
                    guest_id TEXT NOT NULL REFERENCES guests(id), amount_cents INTEGER NOT NULL,
                    count INTEGER NOT NULL CHECK(count>0), first_ticket INTEGER NOT NULL);
                CREATE TABLE IF NOT EXISTS tickets(
                    number INTEGER PRIMARY KEY, allocation_id TEXT NOT NULL REFERENCES allocations(id),
                    status TEXT NOT NULL CHECK(status IN ('active','void')));
                CREATE TABLE IF NOT EXISTS audit(
                    id INTEGER PRIMARY KEY AUTOINCREMENT, at TEXT NOT NULL, action TEXT NOT NULL,
                    payment_id TEXT, details TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS mail_issues(
                    message_id TEXT PRIMARY KEY, received_at TEXT NOT NULL, error TEXT NOT NULL,
                    source_link TEXT NOT NULL, resolved INTEGER NOT NULL DEFAULT 0);
            """)
            defaults = {"revision": 0, "event_title": "A little luck. A lot of love.",
                        "client_id": "", "trusted_sender": "", "account_suffix": "",
                        "start_time": now_iso(), "delta_link": "", "sync_status": "Not connected",
                        "last_sync": None, "excel_pending": True, "excel_error": "", "last_export": None}
            for key, value in defaults.items():
                db.execute("INSERT INTO settings VALUES (?,?) ON CONFLICT(key) DO NOTHING", (key, json.dumps(value)))
            if not db.execute("SELECT 1 FROM guests LIMIT 1").fetchone():
                self._import_guests(db, Path(source))
            elif not self.template.is_file():
                raise RuntimeError("The saved registration template is missing. Restore it from backup before starting.")
        os.chmod(self.path, 0o600)

    @contextmanager
    def snapshot(self):
        with self.lock:
            yield

    @contextmanager
    def connection(self):
        with self.lock:
            db = sqlite3.connect(self.path, timeout=15)
            db.row_factory = sqlite3.Row
            db.execute("PRAGMA foreign_keys=ON")
            try:
                with db:
                    yield db
            finally:
                db.close()

    def _import_guests(self, db, source):
        if not source.is_file():
            raise RuntimeError(f"Registration workbook not found: {source.name}")
        book = load_workbook(source, read_only=True, data_only=True)
        try:
            sheet = book["Raffle List"]
            expected = ["Name", "Remarks", "Guest Of", "Table No.", "Number of tickets purchased", "Amount Paid", "Email Transaction"]
            headers = [sheet.cell(4, col).value for col in range(2, 9)]
            if [str(x).strip().casefold() for x in headers] != [x.casefold() for x in expected]:
                raise RuntimeError("Raffle List headers differ from the agreed template. Review the workbook before importing.")
            for row, values in enumerate(sheet.iter_rows(min_row=5, min_col=2, max_col=8, values_only=True), start=5):
                if not values[0]:
                    continue
                if any(x not in (None, "", 0) for x in values[4:]):
                    raise RuntimeError("Workbook contains existing raffle purchases; zero-balance import was requested.")
                identity = json.dumps([str(x or "").strip() for x in values[:4]], ensure_ascii=False)
                guest_id = str(uuid.uuid5(uuid.NAMESPACE_URL, "wedding-raffle:" + identity))
                db.execute("INSERT INTO guests VALUES (?,?,?,?,?,?)",
                           (guest_id, *[str(x).strip() if x is not None else "" for x in values[:4]], row))
        finally:
            book.close()
        shutil.copy2(source, self.template)
        self._set(db, "source_sha256", hashlib.sha256(source.read_bytes()).hexdigest())
        self._audit(db, "roster_imported", None, {"count": db.execute("SELECT COUNT(*) FROM guests").fetchone()[0]})

    @staticmethod
    def _set(db, key, value):
        db.execute("INSERT INTO settings VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", (key, json.dumps(value)))

    @staticmethod
    def _audit(db, action, payment_id, details):
        db.execute("INSERT INTO audit(at,action,payment_id,details) VALUES (?,?,?,?)",
                   (now_iso(), action, payment_id, json.dumps(details)))

    def _changed(self, db):
        revision = json.loads(db.execute("SELECT value FROM settings WHERE key='revision'").fetchone()[0]) + 1
        self._set(db, "revision", revision)
        self._set(db, "updated_at", now_iso())
        self._set(db, "excel_pending", True)

    def settings(self):
        with self.connection() as db:
            return {row["key"]: json.loads(row["value"]) for row in db.execute("SELECT * FROM settings")}

    def set_values(self, **values):
        with self.connection() as db:
            for key, value in values.items():
                self._set(db, key, value)

    def guests(self):
        with self.connection() as db:
            return [dict(row) for row in db.execute("""
                SELECT g.*, COALESCE(SUM(CASE WHEN p.status='confirmed' THEN a.amount_cents ELSE 0 END),0) paid_cents,
                  COALESCE(SUM(CASE WHEN p.status='confirmed' THEN a.count ELSE 0 END),0) ticket_count
                FROM guests g LEFT JOIN allocations a ON a.guest_id=g.id
                LEFT JOIN payments p ON p.id=a.payment_id GROUP BY g.id ORDER BY LOWER(g.name),g.source_row
            """)]

    def _duplicates(self, db, payment):
        if not payment["bank_time"]:
            return []
        return [row[0] for row in db.execute("""
            SELECT id FROM payments WHERE id!=? AND amount_cents=? AND bank_time=?
              AND COALESCE(account_suffix,'')=COALESCE(?,'') AND status IN ('pending','confirmed')
        """, (payment["id"], payment["amount_cents"], payment["bank_time"], payment["account_suffix"]))]

    def _payment(self, db, row):
        result = dict(row)
        result["allocations"] = [dict(a) for a in db.execute("""
            SELECT a.*,g.name,g.table_no,g.guest_of FROM allocations a JOIN guests g ON a.guest_id=g.id
            WHERE payment_id=? ORDER BY first_ticket
        """, (row["id"],))]
        result["possible_duplicates"] = self._duplicates(db, row)
        try:
            result["ticket_entitlement"] = ticket_count(row["amount_cents"])
            result["amount_error"] = None
        except RuleError as exc:
            result["ticket_entitlement"] = None
            result["amount_error"] = str(exc)
        return result

    def payments(self):
        with self.connection() as db:
            return [self._payment(db, row) for row in db.execute("SELECT * FROM payments ORDER BY created_at DESC,id DESC")]

    def create_payment(self, *, amount_cents, source_key, source="manual", bank_time=None,
                       account_suffix=None, note="", source_link=""):
        if type(amount_cents) is not int or not 0 < amount_cents <= 100_000_000:
            raise RuleError("Enter a positive amount of at most $1,000,000.")
        with self.connection() as db:
            existing = db.execute("SELECT * FROM payments WHERE source_key=?", (source_key,)).fetchone()
            if existing:
                return self._payment(db, existing), False
            payment_id = "PAY-" + uuid.uuid4().hex[:12].upper()
            db.execute("""INSERT INTO payments(id,source_key,source,amount_cents,bank_time,account_suffix,created_at,note,source_link)
                          VALUES (?,?,?,?,?,?,?,?,?)""",
                       (payment_id, source_key, source, amount_cents, bank_time, account_suffix, now_iso(), note, source_link))
            self._audit(db, "payment_imported", payment_id, {"source": source, "amount_cents": amount_cents})
            self._changed(db)
            return self._payment(db, db.execute("SELECT * FROM payments WHERE id=?", (payment_id,)).fetchone()), True

    def confirm(self, payment_id, allocations, verified, duplicate_ack=False):
        with self.connection() as db:
            payment = db.execute("SELECT * FROM payments WHERE id=?", (payment_id,)).fetchone()
            if not payment:
                raise RuleError("Payment not found.")
            # A retry after a successful commit returns the original result, never new tickets.
            if payment["status"] == "confirmed":
                return self._payment(db, payment)
            if payment["status"] != "pending":
                raise RuleError("Only a pending payment can be confirmed.")
            if verified is not True:
                raise RuleError("Verify the payment and buyers before confirming.")
            if self._duplicates(db, payment) and not duplicate_ack:
                raise RuleError("A similar alert already exists. Check the bank record and acknowledge it is a separate transfer.")
            entitlement = ticket_count(payment["amount_cents"])
            if not allocations or len(allocations) > TICKET_CAPACITY:
                raise RuleError("Add at least one guest allocation.")
            for item in allocations:
                if any(type(item.get(k)) is not int for k in ("amount_cents", "count", "first_ticket")):
                    raise RuleError("Amounts, counts and ticket numbers must be integers.")
                if item["amount_cents"] < 0 or item["count"] <= 0:
                    raise RuleError("Guest amounts cannot be negative and ticket counts must be positive.")
                if not db.execute("SELECT 1 FROM guests WHERE id=?", (item.get("guest_id"),)).fetchone():
                    raise RuleError("Select a registered guest for every allocation.")
            if sum(a["amount_cents"] for a in allocations) != payment["amount_cents"]:
                raise RuleError("Guest amounts must add up to the full payment.")
            if sum(a["count"] for a in allocations) != entitlement:
                raise RuleError(f"Allocate exactly {entitlement} tickets for this payment.")
            used = {r[0] for r in db.execute("SELECT number FROM tickets")}
            new_numbers = set()
            for item in allocations:
                first, last = item["first_ticket"], item["first_ticket"] + item["count"] - 1
                if first < FIRST_TICKET or last > LAST_TICKET:
                    raise RuleError(f"Ticket ranges must be within {FIRST_TICKET}–{LAST_TICKET}.")
                numbers = set(range(first, last + 1))
                if numbers & (used | new_numbers):
                    raise RuleError("This ticket range overlaps issued or voided tickets. Check the physical tickets.")
                new_numbers |= numbers
            if len(used) + len(new_numbers) > TICKET_CAPACITY:
                raise RuleError("There are not enough unissued tickets.")
            for item in allocations:
                allocation_id = str(uuid.uuid4())
                db.execute("INSERT INTO allocations VALUES (?,?,?,?,?,?)",
                           (allocation_id, payment_id, item["guest_id"], item["amount_cents"], item["count"], item["first_ticket"]))
                db.executemany("INSERT INTO tickets VALUES (?,?,'active')",
                               [(number, allocation_id) for number in range(item["first_ticket"], item["first_ticket"] + item["count"])])
            db.execute("UPDATE payments SET status='confirmed',confirmed_at=? WHERE id=?", (now_iso(), payment_id))
            self._audit(db, "payment_confirmed", payment_id, {"allocations": allocations, "duplicate_ack": duplicate_ack})
            self._changed(db)
            return self._payment(db, db.execute("SELECT * FROM payments WHERE id=?", (payment_id,)).fetchone())

    def ignore(self, payment_id, reason):
        with self.connection() as db:
            row = db.execute("SELECT status FROM payments WHERE id=?", (payment_id,)).fetchone()
            if not row or row[0] != "pending":
                raise RuleError("Only a pending payment can be ignored.")
            db.execute("UPDATE payments SET status='ignored' WHERE id=?", (payment_id,))
            self._audit(db, "payment_ignored", payment_id, {"reason": reason})
            self._changed(db)

    def reverse(self, payment_id, reason):
        with self.connection() as db:
            row = db.execute("SELECT status FROM payments WHERE id=?", (payment_id,)).fetchone()
            if not row:
                raise RuleError("Payment not found.")
            if row[0] == "reversed":
                return
            if row[0] != "confirmed":
                raise RuleError("Only confirmed payments can be reversed.")
            db.execute("UPDATE payments SET status='reversed' WHERE id=?", (payment_id,))
            db.execute("UPDATE tickets SET status='void' WHERE allocation_id IN (SELECT id FROM allocations WHERE payment_id=?)", (payment_id,))
            self._audit(db, "payment_reversed", payment_id, {"reason": reason, "physical_stubs_must_be_removed": True})
            self._changed(db)

    def reassign(self, payment_id, allocation_id, guest_id, reason):
        with self.connection() as db:
            row = db.execute("SELECT a.*,p.status FROM allocations a JOIN payments p ON a.payment_id=p.id WHERE a.id=? AND p.id=?",
                             (allocation_id, payment_id)).fetchone()
            if not row or row["status"] != "confirmed":
                raise RuleError("Only a confirmed guest allocation can be reassigned.")
            if not db.execute("SELECT 1 FROM guests WHERE id=?", (guest_id,)).fetchone():
                raise RuleError("Select a registered guest.")
            db.execute("UPDATE allocations SET guest_id=? WHERE id=?", (guest_id, allocation_id))
            self._audit(db, "guest_reassigned", payment_id, {"allocation_id": allocation_id, "from": row["guest_id"], "to": guest_id, "reason": reason})
            self._changed(db)

    def summary(self):
        with self.connection() as db:
            pot = db.execute("SELECT COALESCE(SUM(amount_cents),0) FROM payments WHERE status='confirmed'").fetchone()[0]
            sold = db.execute("SELECT COUNT(*) FROM tickets WHERE status='active'").fetchone()[0]
            voided = db.execute("SELECT COUNT(*) FROM tickets WHERE status='void'").fetchone()[0]
            config = {r["key"]: json.loads(r["value"]) for r in db.execute("SELECT * FROM settings")}
            # Deliberately contains no guest, bank, mailbox, source-link or payment identifiers.
            return {"pot_cents": pot, "prizes_cents": prize_cents(pot), "sold": sold, "voided": voided,
                    "remaining": TICKET_CAPACITY - sold - voided, "capacity": TICKET_CAPACITY,
                    "revision": config["revision"], "updated_at": config.get("updated_at"),
                    "event_title": config["event_title"]}

    def audit(self):
        with self.connection() as db:
            return [dict(row) for row in db.execute("SELECT * FROM audit ORDER BY id DESC")]

    def mail_issue(self, message_id, error, source_link=""):
        with self.connection() as db:
            db.execute("INSERT INTO mail_issues(message_id,received_at,error,source_link) VALUES (?,?,?,?) ON CONFLICT(message_id) DO NOTHING",
                       (message_id, now_iso(), error, source_link))

    def mail_issues(self):
        with self.connection() as db:
            return [dict(row) for row in db.execute("SELECT * FROM mail_issues WHERE resolved=0 ORDER BY received_at DESC")]

    def resolve_mail_issue(self, message_id, reason):
        with self.connection() as db:
            if not db.execute("SELECT 1 FROM mail_issues WHERE message_id=?", (message_id,)).fetchone():
                raise RuleError("Alert issue not found.")
            db.execute("UPDATE mail_issues SET resolved=1 WHERE message_id=?", (message_id,))
            self._audit(db, "alert_issue_resolved", None, {"message_id": message_id, "reason": reason})

    def backup(self):
        destination = self.root / "backups"
        destination.mkdir(exist_ok=True, mode=0o700)
        with self.connection() as db:
            path = destination / (now_iso().replace(":", "-") + ".sqlite3")
            target = sqlite3.connect(path)
            try:
                db.backup(target)
            finally:
                target.close()
            os.chmod(path, 0o600)
        for old in sorted(destination.glob("*.sqlite3"))[:-20]:
            old.unlink()
        return path
