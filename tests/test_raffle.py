import copy
import hashlib
import json
import sqlite3
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from uuid import uuid4

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.app import BASE, create_app
from backend.domain import FIRST_TICKET, LAST_TICKET, RuleError, parse_alert, prize_cents, ticket_count
from backend.excel import ExcelExporter
from backend.outlook import GRAPH, Outlook
from backend.store import Store
from backend.reports import audit_reports
from backend.excel import safe_text

SOURCE = BASE / "Registration and Raffle List.xlsx"
ALERT = "You have received SGD 50.00 in your PayNow-linked account ending 3671 on 04-OCT-2026 11:07AM."


@pytest.fixture
def store(tmp_path):
    return Store(tmp_path / "data", SOURCE)


def payment(store, cents=5000, key=None, **kwargs):
    return store.create_payment(amount_cents=cents, source_key=key or str(uuid4()), **kwargs)[0]


def allocation(store, cents=5000, count=15, first=FIRST_TICKET, guest=None):
    return {"guest_id": guest or store.guests()[0]["id"], "amount_cents": cents, "count": count, "first_ticket": first}


@pytest.mark.parametrize("cents,count", [(500,1),(2000,5),(5000,15),(10000,30),(24000,70),(2500,6),(4500,11),(9500,26)])
def test_packages(cents, count):
    assert ticket_count(cents) == count


def test_packages_are_optimal_and_do_not_round():
    for units in range(1, 101):
        best = max(15*f + 5*t + s for f in range(units//10+1) for t in range(units//4+1)
                   for s in [units-10*f-4*t] if s >= 0)
        assert ticket_count(units*500) == best
    for amount in [0,-500,501,24001,5.0,True]:
        with pytest.raises(RuleError):
            ticket_count(amount)
    assert sum(ticket_count(500) for _ in range(4)) == 4


@pytest.mark.parametrize("pot,expected", [(0,[0]*5),(80000,[25000,25000,15000,10000,5000]),
    (160000,[50000,50000,30000,20000,10000]),(240000,[130000,50000,30000,20000,10000])])
def test_prize_examples(pot, expected):
    assert prize_cents(pot) == expected


def test_prizes_reconcile_every_cent_and_keep_ranking():
    for pot in [1,2,3,15,16,17,109999,110000,110001,159999,160001,*range(0,170000,137)]:
        prizes = prize_cents(pot)
        assert sum(prizes) == pot
        assert all(type(p) is int and p >= 0 for p in prizes)
        assert prizes == sorted(prizes, reverse=True)


def test_roster_import_identity_and_restart(store):
    guests = store.guests()
    assert len(guests) == 251
    assert len({g['id'] for g in guests}) == 251
    assert len({g['name'].strip().casefold() for g in guests}) == 250
    assert store.summary()['pot_cents'] == 0
    restarted = Store(store.root, SOURCE)
    assert [g['id'] for g in restarted.guests()] == [g['id'] for g in guests]


def test_group_confirmation_and_retry_are_atomic(store):
    p = payment(store)
    guests = store.guests()
    rows = [allocation(store,2500,8),allocation(store,2500,7,FIRST_TICKET+8,guests[1]['id'])]
    store.confirm(p['id'], rows, True)
    store.confirm(p['id'], rows, True)
    assert store.summary()['pot_cents'] == 5000
    assert store.summary()['sold'] == 15
    assert sum(g['paid_cents'] for g in store.guests()) == 5000
    assert sum(g['ticket_count'] for g in store.guests()) == 15
    assert sum(a['action']=='payment_confirmed' for a in store.audit()) == 1


def test_display_acknowledgements_only_include_new_confirmed_buyers(store):
    baseline = store.display_celebrations()
    assert baseline['contributions'] == []
    guests = store.guests()
    p = payment(store, note='Private bank note', account_suffix='3671')
    rows = [allocation(store,2500,8), allocation(store,2500,7,FIRST_TICKET+8,guests[1]['id'])]
    store.confirm(p['id'], rows, True)
    store.confirm(p['id'], rows, True)
    second = payment(store)
    # The same guest can own multiple physical ranges, but gets one acknowledgement per payment.
    store.confirm(second['id'], [allocation(store,2500,8,FIRST_TICKET+15), allocation(store,2500,7,FIRST_TICKET+23)], True)
    pending = payment(store,500)
    feed = store.display_celebrations(baseline['cursor'])
    assert feed['summary']['pot_cents'] == 10000
    assert set(feed) == {'summary', 'cursor', 'contributions'}
    assert [event['amount_cents'] for event in feed['contributions']] == [5000,5000]
    assert [event['guest_names'] for event in feed['contributions']] == [[guests[0]['name'],guests[1]['name']],[guests[0]['name']]]
    assert all(set(event) == {'id','amount_cents','guest_names'} for event in feed['contributions'])
    for secret in [p['id'],pending['id'],guests[0]['id'],'Private bank note','3671']:
        assert secret not in json.dumps(feed)
    assert store.display_celebrations()['contributions'] == []
    assert store.display_celebrations(feed['cursor'])['contributions'] == []
    assert store.display_celebrations(feed['cursor']+1000)['contributions'] == []
    store.update_guest(guests[0]['id'], name='Corrected name', guest_of='Private relationship', table_no=12)
    assert store.display_celebrations(feed['cursor'])['contributions'] == []
    store.reverse(p['id'],'Refund')
    recovered = store.display_celebrations(baseline['cursor'])
    assert recovered['summary']['pot_cents'] == 5000
    assert [event['amount_cents'] for event in recovered['contributions']] == [5000]


@pytest.mark.parametrize('change', ['amount','count','guest','overlap','low','high','unverified'])
def test_invalid_allocations_never_partially_commit(store, change):
    p=payment(store)
    rows=[allocation(store,2500,8),allocation(store,2500,7,FIRST_TICKET+8)]
    if change=='amount': rows[1]['amount_cents']=2000
    if change=='count': rows[1]['count']=6
    if change=='guest': rows[1]['guest_id']='missing'
    if change=='overlap': rows[1]['first_ticket']=FIRST_TICKET+7
    if change=='low': rows[0]['first_ticket']=FIRST_TICKET-1
    if change=='high': rows[1]['first_ticket']=LAST_TICKET
    with pytest.raises(RuleError): store.confirm(p['id'],rows,change!='unverified')
    assert store.summary()['pot_cents']==store.summary()['sold']==0
    assert store.payments()[0]['allocations']==[]


def test_concurrent_retries_and_conflicting_sales(store):
    p=payment(store,500)
    rows=[allocation(store,500,1)]
    with ThreadPoolExecutor(max_workers=6) as pool:
        results=list(pool.map(lambda _:store.confirm(p['id'],rows,True),range(6)))
    assert len({r['id'] for r in results})==1
    assert store.summary()['sold']==1
    p2=payment(store,500)
    with pytest.raises(RuleError):store.confirm(p2['id'],rows,True)
    assert store.summary()['pot_cents']==500


def test_boundaries_capacity_and_voided_inventory(store):
    p=payment(store,500000)
    store.confirm(p['id'],[allocation(store,500000,1500)],True)
    assert store.summary()['remaining']==0
    assert store.payments()[0]['allocations'][0]['first_ticket']+1499==LAST_TICKET
    extra=payment(store,500)
    with pytest.raises(RuleError):store.confirm(extra['id'],[allocation(store,500,1)],True)
    store.reverse(p['id'],'Event cancellation')
    store.reverse(p['id'],'Retry')
    assert store.summary()['pot_cents']==0
    assert store.summary()['voided']==1500
    assert store.summary()['remaining']==0
    with pytest.raises(RuleError):store.confirm(extra['id'],[allocation(store,500,1)],True)


def test_guest_correction_preserves_tickets_and_audit(store):
    p=payment(store,500)
    store.confirm(p['id'],[allocation(store,500,1, LAST_TICKET)],True)
    a=store.payments()[0]['allocations'][0]
    next_guest=store.guests()[1]['id']
    store.reassign(p['id'],a['id'],next_guest,'Wrong guest selected')
    assert next(g for g in store.guests() if g['id']==next_guest)['ticket_count']==1
    assert store.summary()['pot_cents']==500 and store.summary()['remaining']==1499
    assert store.audit()[0]['action']=='guest_reassigned'


def test_email_parser_and_rejections():
    result=parse_alert('<div>'+ALERT.replace('50.00','1,240.00')+'</div>')
    assert result['amount_cents']==124000
    assert result['bank_time']=='2026-10-04T11:07:00+08:00'
    assert result['account_suffix']=='3671'
    for body in ['You have paid SGD 50.00',ALERT.replace('50.00','5,0.00'),ALERT.replace('04-OCT','99-OCT')]:
        with pytest.raises(RuleError):parse_alert(body)


def message(message_id='message-1', body=ALERT):
    return {'id':message_id,'subject':'UOB-PayNow transfer received','from':{'emailAddress':{'address':'bank@example.test'}},
            'body':{'content':body},'receivedDateTime':'2026-10-04T03:07:10Z','webLink':'https://outlook.live.com/mail/0/inbox/id/example'}


def configured_outlook(store):
    store.set_values(trusted_sender='bank@example.test',account_suffix='3671',start_time='2026-10-04T00:00:00Z',
                     mail_enabled=True,mail_account_id='chosen-account')
    outlook=Outlook(store)
    return outlook


def test_email_dedup_does_not_merge_equal_transfers(store):
    mail=configured_outlook(store)
    assert mail.ingest(message())
    assert not mail.ingest(message())
    assert mail.ingest(message('message-2'))
    assert len(store.payments())==2
    p=store.payments()[0]
    assert len(p['possible_duplicates'])==1
    with pytest.raises(RuleError):store.confirm(p['id'],[allocation(store)],True)
    store.confirm(p['id'],[allocation(store)],True,True)
    assert not mail.ingest({'id':'message-2','@removed':{}})
    assert store.summary()['pot_cents']==5000


def test_email_filtering_and_unparseable_alerts(store):
    mail=configured_outlook(store)
    for changed in [dict(message(),subject='Other transfer'),dict(message(),receivedDateTime='2026-10-03T00:00:00Z'),
                    message(body=ALERT.replace('3671','9999')),dict(message(),**{'from':{'emailAddress':{'address':'someone@example.test'}}})]:
        assert not mail.ingest(changed)
    assert not mail.ingest(message(body='Unexpected bank template'))
    assert len(store.mail_issues())==1
    assert not store.payments()


class FakeMSAL:
    def get_accounts(self):
        return [{'home_account_id':'wrong-account'},{'home_account_id':'chosen-account'}]
    def acquire_token_silent(self, scopes, account):
        assert account['home_account_id']=='chosen-account'
        return {'access_token':'synthetic-test-token'}
    def remove_account(self, account):pass


def mock_graph(monkeypatch, outlook, handler):
    real_client=httpx.Client
    monkeypatch.setattr(outlook,'client',lambda:FakeMSAL())
    monkeypatch.setattr('backend.outlook.httpx.Client',lambda **kwargs:real_client(transport=httpx.MockTransport(handler),**kwargs))


def test_delta_checkpoint_and_restart_after_interrupted_page(store,monkeypatch):
    mail=configured_outlook(store)
    next_url=GRAPH+'/v1.0/me/mailFolders/inbox/messages/delta?skip=next'
    done_url=GRAPH+'/v1.0/me/mailFolders/inbox/messages/delta?delta=done'
    def handler(request):
        assert 'ImmutableId' in request.headers['Prefer']
        if str(request.url)==next_url:raise httpx.ConnectError('offline')
        return httpx.Response(200,json={'value':[message()],'@odata.nextLink':next_url})
    mock_graph(monkeypatch,mail,handler)
    mail.poll()
    assert store.settings()['delta_link']==next_url
    assert len(store.payments())==1
    monkeypatch.undo()
    resumed=Outlook(store)
    def resume(request):
        assert str(request.url)==next_url
        return httpx.Response(200,json={'value':[message(),message('message-2')],'@odata.deltaLink':done_url})
    mock_graph(monkeypatch,resumed,resume)
    resumed.poll()
    assert len(store.payments())==2
    assert store.settings()['delta_link']==done_url


@pytest.mark.parametrize('status,code',[(410,'Gone'),(400,'syncStateNotFound'),(404,'SyncStateNotFound')])
def test_expired_delta_restarts_without_losing_confirmed_sales(store,monkeypatch,status,code):
    mail=configured_outlook(store)
    store.set_values(delta_link=GRAPH+'/v1.0/me/messages/delta?old=1')
    mock_graph(monkeypatch,mail,lambda request:httpx.Response(status,json={'error':{'code':code}}))
    mail.poll()
    assert store.settings()['delta_link']==''
    assert store.settings()['mail_enabled']


def test_rate_limit_and_expired_login_are_visible(store,monkeypatch):
    mail=configured_outlook(store)
    mock_graph(monkeypatch,mail,lambda request:httpx.Response(429,headers={'Retry-After':'120'},json={}))
    mail.poll()
    assert '120s' in store.settings()['sync_status']
    assert mail._retry_at>0
    mail._retry_at=0
    monkeypatch.setattr(mail,'client',lambda:type('Expired',(),{'get_accounts':lambda self:[]})())
    mail.poll()
    assert not store.settings()['mail_enabled']
    assert 'expired' in store.settings()['sync_status']


def test_disconnect_during_response_cannot_restore_old_cursor(store,monkeypatch):
    mail=configured_outlook(store)
    def handler(request):
        mail.disconnect()
        store.set_values(delta_link='',sync_status='New settings')
        return httpx.Response(200,json={'value':[message()],'@odata.deltaLink':GRAPH+'/v1.0/me/messages/delta?old=1'})
    mock_graph(monkeypatch,mail,handler)
    mail.poll()
    assert store.settings()['delta_link']==''
    assert store.settings()['sync_status']=='New settings'
    assert not store.payments()


def test_interactive_login_saves_the_chosen_mailbox_not_first_cached_account(store,monkeypatch):
    mail=configured_outlook(store)
    store.set_values(client_id=str(uuid4()),delta_link=GRAPH+'/v1.0/me/messages/delta?old=1')
    class TwoAccounts:
        def acquire_token_interactive(self,**kwargs):
            return {'access_token':'synthetic','id_token_claims':{'preferred_username':'second@example.test'}}
        def get_accounts(self):
            return [{'home_account_id':'first','username':'first@example.test'},
                    {'home_account_id':'second','username':'second@example.test'}]
    class ImmediateThread:
        def __init__(self,target,**kwargs):self.target=target
        def start(self):self.target()
    monkeypatch.setattr(mail,'client',lambda:TwoAccounts())
    monkeypatch.setattr('backend.outlook.threading.Thread',ImmediateThread)
    mail.start_login()
    config=store.settings()
    assert config['mail_account_id']=='second'
    assert config['mail_username']=='second@example.test'
    assert config['mail_enabled'] and config['delta_link']==''


def test_excel_preserves_source_and_other_sheets_and_retries_lock(store,tmp_path):
    source_hash=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    output=tmp_path/'live.xlsx'
    export=ExcelExporter(store,output)
    assert export.sync()
    lock=output.with_name('~$'+output.name)
    lock.touch()
    p=payment(store,500,note='=WEBSERVICE("unsafe")')
    a=allocation(store,500,1)
    store.confirm(p['id'],[a],True)
    assert not export.sync()
    assert store.settings()['excel_pending'] and store.summary()['pot_cents']==500
    lock.unlink()
    assert export.sync()
    assert not store.settings()['excel_pending']
    original=load_workbook(SOURCE)
    result=load_workbook(output)
    try:
        for name in ['Attendance List','Bridemaid Boxes']:
            assert list(original[name].values)==list(result[name].values)
            assert len(original[name]._images)==len(result[name]._images)
        guest=next(g for g in store.guests() if g['id']==a['guest_id'])
        assert result['Raffle List'].cell(guest['source_row'],6).value==1
        assert result['Raffle List'].cell(guest['source_row'],7).value==5
        assert result['Payment Ledger']['M2'].value.startswith("'=WEBSERVICE")
        assert result['Ticket Register']['A2'].value==FIRST_TICKET
        assert result['Ticket Register']['A1501'].value==LAST_TICKET
        assert result['Raffle List']['K8'].value.startswith('=SUM(')
    finally:original.close();result.close()
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
    store.backup()
    with sqlite3.connect(next((store.root/'backups').glob('*.sqlite3'))) as db:
        assert db.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
        assert db.execute("SELECT SUM(amount_cents) FROM payments WHERE status='confirmed'").fetchone()[0]==500


def test_api_security_idempotency_and_public_privacy(tmp_path):
    app=create_app(data_dir=tmp_path/'data',source=SOURCE,output=tmp_path/'live.xlsx',run_workers=False)
    with TestClient(app) as client:
        assert client.get('/api/state').status_code==403
        assert client.get('/api/session',headers={'Origin':'https://untrusted.example'}).status_code==403
        assert client.get('/api/session',headers={'Host':'untrusted.example'}).status_code==400
        headers={'X-Raffle-Token':client.get('/api/session').json()['token']}
        state=client.get('/api/state',headers=headers).json()
        assert len(state['guests'])==251
        baseline=client.get('/api/display/celebrations').json()
        assert baseline['contributions']==[]
        assert client.get('/api/display/celebrations?after=-1').status_code==422
        assert client.get('/api/display/celebrations?after=9223372036854775808').status_code==422
        request={'request_id':str(uuid4()),'amount_cents':5000}
        p=client.post('/api/payments',json=request,headers=headers).json()['payment']
        assert not client.post('/api/payments',json=request,headers=headers).json()['created']
        body={'verified':True,'allocations':[allocation(app.state.store)]}
        assert client.post(f"/api/payments/{p['id']}/confirm",json=body,headers=headers).status_code==200
        assert client.post(f"/api/payments/{p['id']}/confirm",json=body,headers=headers).status_code==200
        public=client.get('/api/display').json()
        assert public['pot_cents']==5000 and public['sold']==15
        assert set(public)=={'pot_cents','prizes_cents','sold','voided','remaining','capacity','revision','updated_at','event_title'}
        for guest in state['guests']:
            assert guest['id'] not in str(public)
        feed=client.get('/api/display/celebrations',params={'after':baseline['cursor']}).json()
        assert feed['summary']==public
        assert feed['contributions'][0]['guest_names']==[app.state.store.guests()[0]['name']]
        assert client.get('/api/display/celebrations',headers={'Origin':'https://untrusted.example'}).status_code==403
        assert client.get('/api/audit').status_code==403
        assert client.get('/api/audit/sheets').status_code==403
        reports=client.get('/api/audit/sheets',headers=headers).json()
        assert [sheet['name'] for sheet in reports['sheets']]==['Payment Ledger','Ticket Register','Raffle Audit']
        assert len(reports['sheets'][1]['rows'])==1500
        assert reports['sheets'][1]['rows'][0][1]=='active'
        assert client.get('/api/excel/download').status_code==403
        assert client.post(f"/api/payments/{p['id']}/reverse",json={'reason':'   '},headers=headers).status_code==400
        assert client.post('/api/payments',json={'request_id':str(uuid4()),'amount_cents':500.1},headers=headers).status_code==422


def test_audit_reports_match_all_three_excel_sheets_after_group_sales_and_corrections(store,tmp_path):
    p=payment(store)
    guests=store.guests()
    store.confirm(p['id'],[allocation(store,2500,8),allocation(store,2500,7,FIRST_TICKET+8,guests[1]['id'])],True)
    a=next(item for item in store.payments() if item['id']==p['id'])['allocations'][0]
    store.reassign(p['id'],a['id'],guests[2]['id'],'Selected wrong guest')
    reversed_payment=payment(store,500)
    store.confirm(reversed_payment['id'],[allocation(store,500,1,LAST_TICKET)],True)
    store.reverse(reversed_payment['id'],'Refund verified; remove stub')
    ignored=payment(store,2000)
    store.ignore(ignored['id'],'Not a raffle payment')
    payment(store,500,note='=WEBSERVICE("unsafe")')
    reports=audit_reports(store)
    ledger=reports['sheets'][0]['rows']
    assert {row[1] for row in ledger}=={'confirmed','pending','ignored','reversed'}
    assert sum(row[3] or 0 for row in ledger if row[0]==p['id'])==5000
    assert reports['sheets'][1]['rows'][0][1:] == ['active',p['id'],guests[2]['id'],guests[2]['name']]
    assert reports['sheets'][1]['rows'][-1][1]=='void'
    export=ExcelExporter(store,tmp_path/'live.xlsx')
    assert export.sync()
    book=load_workbook(export.output,read_only=True)
    try:
        for report in reports['sheets']:
            actual=list(book[report['name']].values)
            assert list(actual[0])==[column['label'] for column in report['columns']]
            expected=[[value/100 if column['type']=='money' and value is not None else safe_text(value)
                       for value,column in zip(row,report['columns'])] for row in report['rows']]
            # Empty strings round-trip as blank Excel cells.
            assert [list(row) for row in actual[1:]]==[[None if value=='' else value for value in row] for row in expected]
    finally:
        book.close()


def test_guest_management_preserves_identity_history_and_export(store, tmp_path):
    guest_id = str(uuid4())
    guest = store.create_guest(guest_id, name="  New guest  ", guest_of="Friends", table_no=12)
    assert guest["name"] == "New guest" and guest["table_no"] == "12"
    assert store.create_guest(guest_id, name="New guest", guest_of="Friends", table_no=12) == guest
    assert len(store.guests()) == 252
    rows = [allocation(store, guest=guest_id)]
    purchase = store.create_purchase(request_id=uuid4(), amount_cents=5000, note="Verified", allocations=rows, verified=True)
    updated = store.update_guest(guest_id, name="=Updated guest", guest_of="Family", table_no=3)
    assert updated["id"] == guest_id and updated["source_row"] == guest["source_row"]
    assert store.payments()[0]["allocations"][0]["name"] == "=Updated guest"
    store.delete_guest(guest_id)
    store.delete_guest(guest_id)
    assert len(store.guests()) == 251
    assert store.summary()["pot_cents"] == 5000 and store.summary()["sold"] == 15
    removed = next(g for g in store.guests(include_deleted=True) if g["id"] == guest_id)
    assert removed["paid_cents"] == 5000 and removed["deleted_at"]
    with pytest.raises(RuleError, match="registered guest"):
        store.create_purchase(request_id=uuid4(), amount_cents=5000, note="", allocations=[{**rows[0], "first_ticket": FIRST_TICKET+15}], verified=True)
    exporter = ExcelExporter(store, tmp_path / "live.xlsx")
    assert exporter.sync()
    book = load_workbook(exporter.output, data_only=False)
    row = guest["source_row"]
    assert [book["Raffle List"].cell(row, col).value for col in range(2, 10)] == ["'=Updated guest", None, "Family", 3, 15, 50, purchase["id"], "Removed"]
    book.close()
    restarted = Store(store.root, SOURCE)
    assert len(restarted.guests()) == 251 and restarted.summary()["sold"] == 15
    actions = [a["action"] for a in store.audit()]
    assert actions.count("guest_added") == actions.count("guest_updated") == actions.count("guest_deleted") == 1


@pytest.mark.parametrize("fields", [{"name":" "}, {"name":"Guest","table_no":1.5}, {"name":"Guest","table_no":0}, {"name":"Guest","table_no":True}])
def test_guest_validation_is_atomic(store, fields):
    before = store.audit()
    with pytest.raises(RuleError):
        store.create_guest(str(uuid4()), **fields)
    assert len(store.guests()) == 251 and store.audit() == before


def test_direct_purchase_rolls_back_and_retries_without_double_counting(store):
    key = uuid4()
    before, summary = store.audit(), store.summary()
    rows = [allocation(store, count=14)]
    with pytest.raises(RuleError, match="exactly 15"):
        store.create_purchase(request_id=key, amount_cents=5000, note="Reference", allocations=rows, verified=True)
    assert store.payments() == [] and store.audit() == before and store.summary() == summary
    rows[0]["count"] = 15
    purchase = store.create_purchase(request_id=key, amount_cents=5000, note="Reference", allocations=rows, verified=True)
    assert store.create_purchase(request_id=key, amount_cents=5000, note="Reference", allocations=rows, verified=True) == purchase
    assert len(store.payments()) == 1 and store.summary()["sold"] == 15
    with pytest.raises(RuleError, match="already saved"):
        store.create_purchase(request_id=key, amount_cents=2000, note="Reference", allocations=rows, verified=True)
    assert store.summary()["pot_cents"] == 5000


def test_direct_purchase_cross_instance_retry_is_atomic(store):
    key, rows = uuid4(), [allocation(store)]
    def confirm(_):
        return Store(store.root, SOURCE).create_purchase(request_id=key, amount_cents=5000, note="", allocations=rows, verified=True)
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(confirm, range(4)))
    assert len({p["id"] for p in results}) == 1
    assert store.summary()["sold"] == 15 and len(store.payments()) == 1


def test_guest_and_direct_purchase_api_require_auth_and_validate_input(store):
    # Cloud mode skips export on each request; it still uses exactly the shared validation rules.
    class Auth:
        url = "https://test-project.supabase.co"
        async def authorized(self, header):
            return header == "Bearer admin"
    app = create_app(cloud=True, store=store, auth=Auth())
    with TestClient(app) as client:
        for method, path in [("post", "/api/guests"), ("put", "/api/guests/missing"), ("delete", "/api/guests/missing"), ("post", "/api/purchases")]:
            assert getattr(client, method)(path).status_code == 401
        headers = {"Authorization": "Bearer admin"}
        guest = {"request_id":str(uuid4()), "name":"New API guest", "guest_of":"Bride", "table_no":7}
        assert client.post("/api/guests", headers=headers, json={**guest, "table_no":7.5}).status_code == 422
        response = client.post("/api/guests", headers=headers, json=guest)
        assert response.status_code == 200
        guest_id = response.json()["id"]
        purchase = {"request_id":str(uuid4()), "amount_cents":5000, "verified":True, "allocations":[allocation(store, guest=guest_id)]}
        result = client.post("/api/purchases", headers=headers, json=purchase)
        assert result.status_code == 200 and result.json()["status"] == "confirmed"
        assert client.post("/api/purchases", headers=headers, json=purchase).json()["id"] == result.json()["id"]
        assert client.put(f"/api/guests/{guest_id}", headers=headers, json={"name":"Updated API guest", "guest_of":"Groom", "table_no":None}).status_code == 200
        assert client.delete(f"/api/guests/{guest_id}", headers=headers).status_code == 200
        assert len(client.get("/api/state", headers=headers).json()["guests"]) == 251
        assert client.get("/api/display").json()["pot_cents"] == 5000
